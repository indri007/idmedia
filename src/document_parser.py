from pathlib import Path
from io import BytesIO
import pymupdf
from docx import Document

MAX_PAGES = 175
MAX_FILE_BYTES = 50 * 1024 * 1024

def extract_document(filename: str, data: bytes):
    suffix = Path(filename).suffix.lower()
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Ukuran file maksimal 50 MB.")

    if suffix == ".pdf":
        pdf = pymupdf.open(stream=data, filetype="pdf")
        try:
            if len(pdf) > MAX_PAGES:
                raise ValueError("PDF melebihi batas 175 halaman.")
            pages = [
                {"location": f"Halaman {i + 1}", "text": page.get_text()}
                for i, page in enumerate(pdf)
            ]
            if not any(p["text"].strip() for p in pages):
                raise ValueError("PDF tidak memiliki teks yang dapat diekstrak. OCR diperlukan untuk PDF hasil scan.")
            return {
                "pages": pages,
                "page_count": len(pdf),
                "tables": None,
                "images": sum(len(p.get_images()) for p in pdf),
                "format": "pdf",
            }
        finally:
            pdf.close()

    if suffix == ".docx":
        doc = Document(BytesIO(data))
        pages = []
        for i, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                pages.append({
                    "location": f"Paragraf {i}",
                    "text": paragraph.text,
                    "style": paragraph.style.name if paragraph.style else "",
                })

        tables = []
        for ti, table in enumerate(doc.tables, 1):
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            tables.append({"table": ti, "rows": rows})

        images = len(doc.inline_shapes)
        if not pages and not tables:
            raise ValueError("Tidak ditemukan teks atau tabel di DOCX.")
        return {
            "pages": pages,
            "page_count": None,
            "tables": tables,
            "images": images,
            "format": "docx",
        }

    raise ValueError("Format tidak didukung. Gunakan PDF atau DOCX.")

def extract_text(filename: str, data: bytes):
    """Kompatibilitas dengan versi MVP sebelumnya."""
    result = extract_document(filename, data)
    return "\n".join(p["text"] for p in result["pages"]), result["page_count"]
