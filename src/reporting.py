from io import BytesIO
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

def make_docx_report(filename, findings, references):
    doc = Document()
    doc.add_heading("Scientific Manuscript QA — Laporan Audit", 0)
    doc.add_paragraph(f"Dokumen: {filename}")
    doc.add_paragraph(
        "Laporan otomatis untuk membantu revisi. Temuan harus diverifikasi oleh penulis."
    )
    doc.add_heading("Temuan", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Prioritas", "Kategori", "Lokasi", "Temuan", "Saran"]):
        cell.text = value
    for f in findings:
        cells = table.add_row().cells
        for cell, value in zip(cells, [
            f["prioritas"], f["kategori"], f["lokasi"], f["temuan"], f["saran"]
        ]):
            cell.text = str(value)
    doc.add_heading("Pemeriksaan metadata referensi", level=1)
    for ref in references:
        doc.add_paragraph(
            f"{ref.get('provider', '')}: {ref.get('reference', '')}\n"
            f"Match: {ref.get('match')} | Judul kandidat: {ref.get('title', '')} | "
            f"DOI: {ref.get('doi', '')}"
        )
    out = BytesIO()
    doc.save(out)
    return out.getvalue()

def make_pdf_report(filename, findings, references):
    out = BytesIO()
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Scientific Manuscript QA — Laporan Audit", styles["Title"]),
        Paragraph("Dokumen: " + escape(filename), styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Temuan otomatis — verifikasi sebelum melakukan revisi", styles["Heading2"]),
    ]
    rows = [["Prioritas", "Kategori", "Lokasi", "Temuan", "Saran"]]
    for f in findings:
        rows.append([
            f["prioritas"], f["kategori"], f["lokasi"],
            Paragraph(escape(f["temuan"]), styles["BodyText"]),
            Paragraph(escape(f["saran"]), styles["BodyText"]),
        ])
    table = Table(rows, repeatRows=1, colWidths=[55, 65, 65, 145, 165])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph("Metadata referensi (kandidat, bukan validasi final)", styles["Heading2"]))
    for ref in references:
        line = (
            f"{ref.get('provider', '')}: {ref.get('reference', '')} — "
            f"Match: {ref.get('match')} — DOI: {ref.get('doi', '')}"
        )
        story.append(Paragraph(escape(line), styles["BodyText"]))
        story.append(Spacer(1, 5))
    SimpleDocTemplate(out, pagesize=A4, rightMargin=30, leftMargin=30).build(story)
    return out.getvalue()
